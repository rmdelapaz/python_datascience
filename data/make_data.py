"""Generate the Pinecrest Outfitters course data.

Pinecrest Outfitters is a FICTIONAL outdoor-gear retailer (an online shop plus
stores in Boise, Bend, Flagstaff and Asheville). Every person, product and
number in these files is invented for the Python Data Science course.

    python data/make_data.py            # write everything into data/
    python data/make_data.py --out DIR  # write somewhere else

Needs only the standard library, numpy and pandas. The Excel workbook also
needs openpyxl (the engine pandas uses for .xlsx); without it that one file is
skipped with a message.

Reproducibility: the generator is seeded, so running it twice gives
byte-identical CSV, JSON, XLSX and ZIP files. pinecrest.db is byte-identical
when built with the same SQLite library version (SQLite stamps its version
number into the file header); its contents are identical everywhere.

Fictional data written for this course. (c) 2026 Ray de la Paz.
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sqlite3
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

SEED = 20260112
START, END = "2024-01-01", "2025-12-31"
STORES = ["Boise", "Bend", "Flagstaff", "Asheville"]
STORE_STATE = {"Boise": "ID", "Bend": "OR", "Flagstaff": "AZ", "Asheville": "NC"}
FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)


def rng_for(n: int) -> np.random.Generator:
    """An independent, reproducible random stream per dataset."""
    return np.random.default_rng([SEED, n])


def to_csv(df: pd.DataFrame, path: Path) -> None:
    df.to_csv(path, index=False, lineterminator="\n", encoding="utf-8")


# --------------------------------------------------------------------------
# Names (all invented combinations)
# --------------------------------------------------------------------------
FIRST = """Aaliyah Abel Adaeze Adrian Aiko Alana Alejandro Amara Amir Ana Andre Anika Anton Aria Arjun Asha
Beatriz Ben Bianca Bodhi Brielle Bruno Caleb Camila Carmen Cedric Chidi Chloe Clara Cole Dalia Damon Dara Darius
Deepa Delia Desmond Diego Dmitri Elan Elena Eli Elise Emeka Esme Ethan Fatima Felix Fiona Gabriel Gemma Gideon
Grace Hana Hassan Hazel Hector Helga Hugo Idris Imani Ingrid Isaac Isla Ivan Jada Jalen Jasper Javier Jin
Joanna Jonah Josefina Jude Kai Kalani Kareem Katya Keanu Keiko Kenji Kiran Lars Layla Leila Leo Liam Lidia
Linh Lola Lucia Luka Mabel Malik Marco Maren Mateo Maya Mei Micah Mila Milan Nadia Naoko Nia Nico Nikhil Noor
Odessa Olga Omar Oren Paloma Pavel Petra Priya Quinn Rafael Rania Ravi Reese Renata Rhea Rohan Rosa Ruben
Sadie Samir Sana Santiago Selma Sienna Silas Soren Tamar Tariq Tess Theo Tomas Uma Valentina Vera Vikram
Wes Willa Xavier Yara Yusuf Zane Zara Zoe""".split()
LAST = """Abara Acosta Adeyemi Aguilar Alder Almeida Amato Anand Arce Ashby Baek Bakshi Banerjee Barlow Becker
Bello Bergstrom Bianchi Blackwood Boone Brandt Calloway Campos Castellan Cha Chaudhary Cho Clarke Cordova
Dalton Danso Delgado Demir Desai Diaz Dubois Eastman Ekwueme Ellery Espinoza Farrow Faulkner Ferrand Fontaine
Fujita Gallo Garza Getty Gill Goswami Greer Haddad Halvorsen Hartley Hayashi Herrera Holloway Hosseini Ibarra
Idowu Ito Jaramillo Joshi Kaplan Karimi Kato Kessler Khalil Kimura Kowalski Kumar Lachance Laine Larkin
Leclerc Lindqvist Lozano Madrigal Mahlangu Marchetti Marlowe Mbeki Mendez Mistry Moreau Mori Nakata Navarro
Ndiaye Novak Nwosu Oakes Obi Okafor Olsen Ortega Osei Pacheco Paredes Pellegrino Petrov Pham Pike Prescott
Quan Quintero Rahimi Raines Ramos Rasmussen Reyes Rinaldi Rojas Rowan Sabry Salazar Sato Schofield Sen Shah
Silva Sokolov Soto Stanek Sundberg Tamura Tanaka Tavares Thackeray Tran Ueda Underhill Vance Varga Vasquez
Villareal Volkov Wainwright Watanabe Whitlock Winslow Xiong Yamada Yilmaz Yoon Zamora Zapata Zeller""".split()

REGION_CITIES = {
    "ID": ["Boise", "Meridian", "Nampa", "Eagle", "Caldwell"],
    "OR": ["Bend", "Redmond", "Sisters", "Prineville", "Eugene", "Portland"],
    "AZ": ["Flagstaff", "Sedona", "Prescott", "Williams", "Phoenix", "Tucson"],
    "NC": ["Asheville", "Hendersonville", "Black Mountain", "Brevard", "Charlotte", "Raleigh"],
}
OTHER_CITIES = [("Seattle", "WA"), ("Spokane", "WA"), ("Denver", "CO"), ("Salt Lake City", "UT"),
                ("Reno", "NV"), ("Sacramento", "CA"), ("Missoula", "MT"), ("Knoxville", "TN"),
                ("Greenville", "SC"), ("Atlanta", "GA"), ("Albuquerque", "NM"), ("Las Vegas", "NV")]
AREA_CODE = {"ID": "208", "OR": "541", "AZ": "928", "NC": "828", "WA": "206", "CO": "303", "UT": "801",
             "NV": "775", "CA": "916", "MT": "406", "TN": "865", "SC": "864", "GA": "404", "NM": "505"}


# --------------------------------------------------------------------------
# Products
# --------------------------------------------------------------------------
BRANDS = ["Pinecrest", "Larkspur Ridge", "Granite Fork", "Wrenfield", "Cedar Notch", "Tallgrass",
          "Juniper Vane", "Quillback"]
MODELS = ["Summit", "Switchback", "Timberline", "Ridgeline", "Cascade", "Basecamp", "Trailhead",
          "Alpenglow", "Skyline", "Backcountry", "Wayfinder", "Cairn", "Talus", "Meadow", "Ember",
          "Drift", "Solstice", "Tundra", "Canyon", "Foothill", "Lookout", "Headwater", "Sawtooth", "Mesa"]
COLORS = ["Slate", "Forest", "Rust", "Sand", "Charcoal", "Glacier Blue", "Moss", "Ember Orange", "Plum", "Stone"]
# category: (sku prefix, [(type, low price, high price, weight kg, max qty per line)])
CATALOG = {
    "Tents": ("TNT", [("1-Person Tent", 150, 330, 1.3, 1), ("2-Person Tent", 180, 480, 1.9, 1),
                      ("3-Person Tent", 240, 560, 2.6, 1), ("4-Person Tent", 280, 620, 3.9, 1),
                      ("6-Person Family Tent", 320, 700, 6.8, 1), ("Ultralight Shelter", 220, 520, 0.8, 1),
                      ("Tarp", 50, 140, 0.6, 2), ("Tent Footprint", 30, 70, 0.3, 2)]),
    "Sleeping": ("SLP", [("20F Sleeping Bag", 120, 380, 1.2, 2), ("30F Sleeping Bag", 90, 300, 1.0, 2),
                         ("0F Sleeping Bag", 220, 520, 1.8, 1), ("Sleeping Pad", 45, 210, 0.6, 2),
                         ("Bag Liner", 25, 75, 0.2, 2), ("Camp Pillow", 18, 55, 0.15, 2)]),
    "Packs": ("PCK", [("22L Daypack", 55, 140, 0.6, 1), ("30L Daypack", 80, 180, 0.9, 1),
                      ("50L Backpack", 160, 320, 1.6, 1), ("65L Backpack", 200, 380, 2.0, 1),
                      ("Hydration Pack", 60, 150, 0.5, 1), ("Duffel", 70, 190, 1.1, 1), ("Dry Bag", 15, 45, 0.1, 3)]),
    "Footwear": ("FTW", [("Hiking Boots", 130, 290, 1.1, 1), ("Trail Runners", 110, 180, 0.6, 1),
                         ("Camp Sandals", 45, 120, 0.5, 1), ("Hiking Socks", 14, 28, 0.08, 6),
                         ("Gaiters", 30, 80, 0.2, 1), ("Insoles", 25, 60, 0.1, 2)]),
    "Apparel": ("APL", [("Rain Jacket", 90, 320, 0.4, 1), ("Down Jacket", 160, 420, 0.4, 1),
                        ("Fleece Pullover", 55, 150, 0.35, 2), ("Base Layer Top", 40, 110, 0.2, 2),
                        ("Hiking Pants", 60, 140, 0.35, 2), ("Trail Shorts", 40, 85, 0.2, 2),
                        ("Sun Hat", 24, 55, 0.1, 2), ("Beanie", 18, 40, 0.08, 3), ("Liner Gloves", 20, 50, 0.06, 2)]),
    "Climbing": ("CLM", [("Climbing Harness", 55, 130, 0.4, 1), ("Climbing Helmet", 60, 140, 0.35, 1),
                         ("Locking Carabiner", 10, 24, 0.07, 6), ("60m Rope", 170, 300, 3.6, 1),
                         ("Chalk Bag", 16, 36, 0.09, 2), ("Belay Device", 22, 110, 0.1, 1),
                         ("Quickdraw Set", 75, 150, 0.6, 2)]),
    "Camp Kitchen": ("KIT", [("Backpacking Stove", 45, 170, 0.3, 1), ("Cookset", 40, 130, 0.5, 1),
                             ("Insulated Mug", 16, 38, 0.2, 4), ("Water Filter", 30, 120, 0.2, 1),
                             ("Water Bottle", 12, 45, 0.2, 4), ("Cooler", 60, 320, 4.5, 1),
                             ("Titanium Spork", 6, 16, 0.02, 6), ("Fuel Canister", 6, 12, 0.23, 6)]),
    "Snow": ("SNW", [("Snowshoes", 140, 300, 1.9, 1), ("Snow Goggles", 50, 190, 0.15, 1),
                     ("Insulated Gloves", 45, 130, 0.25, 2), ("Insulated Bib", 150, 340, 0.8, 1),
                     ("Sled", 30, 90, 2.2, 2), ("Traction Cleats", 25, 75, 0.35, 2), ("Hand Warmers 10-Pack", 9, 18, 0.3, 6)]),
    "Accessories": ("ACC", [("Headlamp", 22, 110, 0.09, 3), ("Lantern", 28, 90, 0.4, 2), ("Compass", 14, 70, 0.05, 2),
                            ("Trekking Poles", 60, 190, 0.5, 1), ("Multitool", 30, 120, 0.2, 1),
                            ("First-Aid Kit", 18, 65, 0.3, 2), ("Sunscreen SPF 50", 9, 18, 0.1, 4),
                            ("Insect Repellent", 8, 15, 0.1, 4), ("Trail Map Set", 12, 30, 0.1, 2)]),
}
# How strongly each category sells in each month (Jan..Dec); 1.0 = average.
CAT_SEASON = {
    "Tents":        [0.4, 0.5, 0.8, 1.1, 1.6, 1.9, 1.8, 1.4, 0.9, 0.6, 0.5, 0.6],
    "Sleeping":     [0.5, 0.6, 0.8, 1.1, 1.5, 1.7, 1.6, 1.3, 1.0, 0.7, 0.6, 0.7],
    "Packs":        [0.7, 0.7, 0.9, 1.1, 1.3, 1.4, 1.4, 1.3, 1.0, 0.8, 0.8, 1.0],
    "Footwear":     [0.8, 0.8, 1.0, 1.2, 1.3, 1.2, 1.1, 1.1, 1.1, 1.0, 0.9, 1.0],
    "Apparel":      [1.1, 1.0, 0.9, 0.9, 0.9, 0.8, 0.8, 0.8, 1.0, 1.2, 1.4, 1.7],
    "Climbing":     [0.6, 0.7, 1.0, 1.3, 1.4, 1.3, 1.1, 1.1, 1.2, 1.1, 0.7, 0.8],
    "Camp Kitchen": [0.6, 0.6, 0.8, 1.0, 1.4, 1.6, 1.6, 1.4, 1.0, 0.7, 0.7, 1.0],
    "Snow":         [2.4, 2.0, 0.9, 0.2, 0.1, 0.05, 0.05, 0.05, 0.2, 0.8, 2.0, 3.0],
    "Accessories":  [0.8, 0.8, 0.9, 1.0, 1.2, 1.3, 1.3, 1.2, 1.0, 0.9, 1.0, 1.3],
}
N_PER_CATEGORY = {"Tents": 24, "Sleeping": 24, "Packs": 26, "Footwear": 28, "Apparel": 36,
                  "Climbing": 24, "Camp Kitchen": 30, "Snow": 24, "Accessories": 32}   # 248 before removals
N_DISCONTINUED, N_NEW = 6, 8
PRICE_CHANGE_DATE = "2025-03-01"


def make_products():
    """Return (catalog incl. discontinued items, public products table)."""
    rng = rng_for(1)
    rows, used = [], set()
    for cat, (prefix, types) in CATALOG.items():
        for i in range(N_PER_CATEGORY[cat]):
            t, lo, hi, wt, maxq = types[i % len(types)]
            while True:
                name = f"{rng.choice(BRANDS)} {rng.choice(MODELS)} {t}"
                if name not in used:
                    used.add(name)
                    break
            dollars = int(rng.uniform(lo, hi))
            price = float(dollars) + (0.95 if dollars >= 20 else 0.5)
            cost = round(price * rng.uniform(0.42, 0.66), 2)
            rows.append({
                "sku": f"{prefix}-{1001 + i}", "name": name, "category": cat, "product_type": t,
                "brand": name.split(f" {t}")[0].rsplit(" ", 1)[0],
                "price": price, "cost": cost,
                "weight_kg": round(wt * rng.uniform(0.8, 1.25), 2),
                "color": str(rng.choice(COLORS)),
                "_maxq": maxq, "_pop": float(rng.lognormal(0, 0.9)) * (60.0 / price) ** 0.55,
            })
    cat_df = pd.DataFrame(rows)
    n = len(cat_df)
    cat_df["launch_date"] = pd.Timestamp("2019-01-01") + pd.to_timedelta(rng.integers(0, 365 * 5, n), unit="D")
    mid = rng.choice(n, 18, replace=False)            # launched during the data period
    cat_df.loc[mid, "launch_date"] = (pd.Timestamp("2024-02-01") + pd.to_timedelta(rng.integers(0, 600, 18), unit="D")).values
    order = rng.permutation(n)
    new_idx, disc_idx = order[:N_NEW], order[N_NEW:N_NEW + N_DISCONTINUED]
    cat_df.loc[new_idx, "launch_date"] = (pd.Timestamp("2025-12-20") + pd.to_timedelta(rng.integers(0, 10, N_NEW), unit="D")).values
    cat_df["_new"] = False
    cat_df.loc[new_idx, "_new"] = True
    cat_df["_disc"] = False
    cat_df.loc[disc_idx, "_disc"] = True
    cat_df.loc[disc_idx, "launch_date"] = pd.Timestamp("2021-06-01")
    # ratings: new products have no reviews yet
    reviews = np.round(rng.lognormal(3.6, 1.1, n) * np.sqrt(cat_df["_pop"])).astype(int)
    reviews[cat_df["_new"].to_numpy()] = 0
    rating = np.clip(rng.normal(4.15, 0.42, n), 2.4, 5.0).round(1)
    cat_df["rating"] = np.where(reviews > 0, rating, np.nan)
    cat_df["review_count"] = reviews
    # 2025-03-01 price increase on about 60% of the catalog
    raised = rng.random(n) < 0.6
    pct = rng.uniform(0.04, 0.08, n)
    old = np.floor(cat_df["price"] * (1 - pct)) + (cat_df["price"] % 1)
    cat_df["_old_price"] = np.where(raised, np.minimum(old, cat_df["price"] - 1).round(2), cat_df["price"])
    # a few missing weights (supplier never sent the spec)
    miss = rng.choice(n, 14, replace=False)
    cat_df.loc[miss, "weight_kg"] = np.nan
    return cat_df


# --------------------------------------------------------------------------
# Weather and daily store sales
# --------------------------------------------------------------------------
CLIMATE = {  # mean high F, seasonal amplitude, diurnal range, wet-day probability by month, rain scale (in)
    "Boise":     (63, 27, 24, [.30, .28, .26, .22, .20, .12, .05, .05, .10, .17, .28, .32], 0.16),
    "Bend":      (59, 23, 28, [.28, .25, .23, .18, .17, .12, .07, .07, .09, .15, .26, .30], 0.15),
    "Flagstaff": (61, 20, 27, [.22, .22, .22, .14, .09, .05, .40, .45, .25, .13, .13, .20], 0.26),
    "Asheville": (67, 19, 21, [.34, .34, .37, .34, .38, .42, .44, .40, .30, .25, .30, .34], 0.36),
}
PROMO_WINDOWS = [("2024-05-24", "2024-05-27"), ("2024-08-30", "2024-09-02"), ("2024-11-29", "2024-12-02"),
                 ("2025-05-23", "2025-05-26"), ("2025-08-29", "2025-09-01"), ("2025-11-28", "2025-12-01")]
CLOSED_DAYS = ["2024-11-28", "2024-12-25", "2025-11-27", "2025-12-25"]          # every store
EXPORT_GAP = {"Bend": ["2024-08-12", "2024-08-13", "2024-08-14"], "Flagstaff": ["2025-01-21", "2025-01-22"]}
STORE_BASE = {  # visitors/day, yearly growth, staff baseline
    "Boise": (420, 0.08, 7.0), "Bend": (300, 0.03, 5.0), "Flagstaff": (250, 0.22, 4.4), "Asheville": (280, -0.06, 5.0)}
DOW_FACTOR = np.array([0.84, 0.82, 0.86, 0.90, 1.08, 1.52, 1.26])     # Mon..Sun


def season_curve(doy: np.ndarray) -> np.ndarray:
    """Outdoor retail year: a summer hump, a December holiday peak, a January-February lull."""
    return (1.0 + 0.32 * np.exp(-((doy - 195) / 48.0) ** 2) + 0.42 * np.exp(-((doy - 349) / 11.0) ** 2)
            - 0.16 * np.exp(-((doy - 38) / 30.0) ** 2))


def make_weather(days: pd.DatetimeIndex) -> pd.DataFrame:
    rng = rng_for(2)
    doy = days.dayofyear.to_numpy()
    frames = []
    for city, (mean, amp, diurnal, wet_p, scale) in CLIMATE.items():
        noise = np.zeros(len(days))
        shocks = rng.normal(0, 5.2, len(days))
        for i in range(1, len(days)):
            noise[i] = 0.68 * noise[i - 1] + shocks[i]
        wet = rng.random(len(days)) < np.array(wet_p)[days.month.to_numpy() - 1]
        precip = np.where(wet, np.maximum(rng.gamma(0.8, scale, len(days)), 0.01), 0.0).round(2)
        high = mean + amp * np.cos(2 * np.pi * (doy - 203) / 365.25) + noise - 4.5 * wet
        warm = (np.cos(2 * np.pi * (doy - 203) / 365.25) + 1) / 2          # 0 in midwinter, 1 in midsummer
        low = high - diurnal * (0.58 + 0.42 * warm) + rng.normal(0, 3.5, len(days)) + 5 * wet
        frames.append(pd.DataFrame({"date": days, "city": city, "high_f": high.round(0).astype(int),
                                    "low_f": np.minimum(low, high - 3).round(0).astype(int), "precip_in": precip}))
    return pd.concat(frames, ignore_index=True)


def make_daily_sales(days: pd.DatetimeIndex, weather: pd.DataFrame) -> pd.DataFrame:
    """Point-of-sale summary: ALL in-store sales, including walk-ins that never appear in orders.csv."""
    rng = rng_for(3)
    doy = days.dayofyear.to_numpy()
    years_in = (np.arange(len(days))) / 365.0
    promo = np.zeros(len(days), dtype=int)
    for a, b in PROMO_WINDOWS:
        promo[(days >= a) & (days <= b)] = 1
    frames = []
    for store in STORES:
        base, growth, staff_base = STORE_BASE[store]
        w = weather[weather["city"] == store]
        rainy = (w["precip_in"].to_numpy() >= 0.10)
        comfort = 1 - 0.0022 * np.abs(w["high_f"].to_numpy() - 72)
        expected = (1 + growth * years_in) * season_curve(doy) * DOW_FACTOR[days.dayofweek.to_numpy()] * (1 + 0.45 * promo)
        visitors = np.round(base * expected * (1 - 0.21 * rainy) * comfort * rng.lognormal(0, 0.075, len(days))).astype(int)
        conv = np.clip(0.235 * (1 + 0.10 * promo) * rng.lognormal(0, 0.05, len(days)), 0.05, 0.6)
        transactions = rng.binomial(visitors, conv)
        units = transactions + rng.poisson(1.12 * transactions)
        unit_price = 41.0 * (1 + 0.10 * np.exp(-((doy - 349) / 30.0) ** 2) + 0.06 * (doy < 60)) * rng.lognormal(0, 0.055, len(days))
        revenue = (units * unit_price).round(2)
        # Staffing is planned from EXPECTED demand (season, weekday, promotions) - not from the day's luck.
        staff = np.maximum(2, np.round(staff_base * expected + rng.normal(0, 0.55, len(days)))).astype(int)
        frames.append(pd.DataFrame({"date": days, "store": store, "visitors": visitors, "transactions": transactions,
                                    "units": units, "revenue": revenue, "staff_on_duty": staff, "promo_day": promo}))
    daily = pd.concat(frames, ignore_index=True).sort_values(["date", "store"], kind="stable").reset_index(drop=True)
    return daily


def punch_gaps(daily: pd.DataFrame) -> pd.DataFrame:
    """Remove closed days / failed exports and blank a few revenue cells (the 'few gaps')."""
    rng = rng_for(4)
    d = daily.copy()
    keep = ~d["date"].isin(pd.to_datetime(CLOSED_DAYS))
    for store, dates in EXPORT_GAP.items():
        keep &= ~((d["store"] == store) & d["date"].isin(pd.to_datetime(dates)))
    d = d[keep].reset_index(drop=True)
    blanks = rng.choice(len(d), 6, replace=False)
    d.loc[blanks, "revenue"] = np.nan
    return d


# --------------------------------------------------------------------------
# Customers, orders, order items
# --------------------------------------------------------------------------
N_CUSTOMERS, N_CLOSED = 5000, 150
ONLINE_PER_DAY = 14.0
MEMBER_SHARE = 0.030      # share of in-store transactions made by signed-in account holders


def make_customers():
    rng = rng_for(5)
    n = N_CUSTOMERS
    first = rng.choice(FIRST, n)
    last = rng.choice(LAST, n)
    region = rng.choice(["ID", "OR", "AZ", "NC", "other"], n, p=[0.22, 0.17, 0.15, 0.17, 0.29])
    city, state = [], []
    for r in region:
        if r == "other":
            c, s = OTHER_CITIES[rng.integers(len(OTHER_CITIES))]
        else:
            cities = REGION_CITIES[r]
            pw = np.array([0.45] + [0.55 / (len(cities) - 1)] * (len(cities) - 1))
            c, s = str(rng.choice(cities, p=pw)), r
        city.append(c)
        state.append(s)
    # two thirds signed up before the data period, the rest during it
    before = rng.random(n) < 0.66
    signup = np.where(before,
                      (pd.Timestamp("2019-01-01") + pd.to_timedelta(rng.integers(0, 1826, n), unit="D")).values,
                      (pd.Timestamp("2024-01-01") + pd.to_timedelta(rng.integers(0, 715, n), unit="D")).values)
    seen: dict[str, int] = {}
    emails = []
    for f, l in zip(first, last):
        base = f"{f}.{l}".lower()
        seen[base] = seen.get(base, 0) + 1
        emails.append(f"{base}{'' if seen[base] == 1 else seen[base]}@example.com")
    phone = [f"({AREA_CODE[s]}) 555-01{rng.integers(0, 100):02d}" for s in state]
    phone = np.where(rng.random(n) < 0.06, None, np.array(phone, dtype=object))
    birth = rng.integers(1948, 2007, n).astype(float)
    birth[rng.random(n) < 0.09] = np.nan
    acq = rng.choice(["organic", "paid_search", "social", "referral", "in_store"], n, p=[0.30, 0.22, 0.16, 0.12, 0.20])
    cust = pd.DataFrame({
        "customer_id": np.arange(10001, 10001 + n), "first_name": first, "last_name": last, "email": emails,
        "phone": phone, "city": city, "state": state, "signup_date": pd.to_datetime(signup),
        "birth_year": birth, "acquisition_channel": acq, "marketing_opt_in": rng.random(n) < 0.58,
    })
    weight = rng.gamma(1.5, 1.0, n)
    weight[rng.random(n) < 0.06] = 0.0            # signed up, never bought
    cust["_w"] = weight
    return cust


def make_orders(cust: pd.DataFrame, catalog: pd.DataFrame, daily_full: pd.DataFrame, days: pd.DatetimeIndex):
    rng = rng_for(6)
    nd = len(days)
    doy = days.dayofyear.to_numpy()
    promo = np.zeros(nd, dtype=int)
    for a, b in PROMO_WINDOWS:
        promo[(days >= a) & (days <= b)] = 1
    black_friday = np.zeros(nd, dtype=bool)
    for a, b in (PROMO_WINDOWS[2], PROMO_WINDOWS[5]):
        black_friday |= (days >= a) & (days <= b)
    # --- how many orders on each day in each stream -------------------------------
    online_dow = np.array([1.08, 1.02, 1.0, 0.98, 0.92, 0.95, 1.10])
    lam = ONLINE_PER_DAY * (1 + 0.16 * np.arange(nd) / 365.0) * (0.75 + 0.25 * season_curve(doy)) \
        * online_dow[days.dayofweek.to_numpy()] * (1 + 0.9 * promo) * (1 + 0.55 * np.exp(-((doy - 349) / 9.0) ** 2))
    counts = {"Online": rng.poisson(lam)}
    for store in STORES:
        tx = daily_full.loc[daily_full["store"] == store, "transactions"].to_numpy()
        tx = np.where(days.isin(pd.to_datetime(CLOSED_DAYS)), 0, tx)
        counts[store] = rng.binomial(tx, MEMBER_SHARE)
    # --- who buys ---------------------------------------------------------------
    signup_day = (cust["signup_date"] - days[0]).dt.days.to_numpy()
    w = cust["_w"].to_numpy()
    state = cust["state"].to_numpy()
    aff = {"Online": np.where(np.isin(state, list(STORE_STATE.values())), 0.38, 1.0)}
    for store in STORES:
        aff[store] = np.where(state == STORE_STATE[store], 1.0, 0.012)
    slots = []
    for stream in ["Online"] + STORES:
        base_p = w * aff[stream]
        for d in range(nd):
            k = counts[stream][d]
            if k == 0:
                continue
            p = np.where(signup_day <= d, base_p, 0.0)
            who = rng.choice(len(cust), size=k, p=p / p.sum())
            slots.extend((d, stream, int(cust["customer_id"].iat[i])) for i in who)
    slots.sort(key=lambda s: (s[0], s[1] != "Online", s[1]))    # by day; stable within a day
    perm = np.concatenate([rng.permutation(np.flatnonzero(np.array([s[0] for s in slots]) == d)) for d in range(nd)])
    slots = [slots[i] for i in perm]
    # --- what they buy ------------------------------------------------------------
    cat_arr = catalog["category"].to_numpy()
    launch_day = (catalog["launch_date"] - days[0]).dt.days.to_numpy()
    pop = catalog["_pop"].to_numpy() * np.where(catalog["_new"], 0.0, 1.0)
    disc_cut = (pd.Timestamp("2024-09-30") - days[0]).days
    price_cut = (pd.Timestamp(PRICE_CHANGE_DATE) - days[0]).days
    month_p = {}
    for ym in sorted({(d.year, d.month) for d in days}):
        first_idx = (pd.Timestamp(year=ym[0], month=ym[1], day=1) - days[0]).days
        season = np.array([CAT_SEASON[c][ym[1] - 1] for c in cat_arr])
        p = pop * season * (launch_day <= first_idx + 27)
        p = np.where(catalog["_disc"] & (first_idx > disc_cut), 0.0, p)
        month_p[ym] = p / p.sum()
    sku = catalog["sku"].to_numpy()
    price_now = np.round(catalog["price"].to_numpy() * 100).astype(int)
    price_old = np.round(catalog["_old_price"].to_numpy() * 100).astype(int)
    maxq = catalog["_maxq"].to_numpy()
    first_seen: set[int] = set()
    plant = rng_for(14)    # own stream: after the 2025-03-01 rise, shoppers buy fewer units of repriced items
    online_slots = np.flatnonzero(np.array([s[1] == "Online" for s in slots]))
    bulk_at = set(rng.choice(online_slots, 12, replace=False).tolist())      # outfitters buying for a group
    bulk_pool = np.flatnonzero((catalog["price"].to_numpy() >= 60) & (catalog["price"].to_numpy() <= 160)
                               & ~catalog["_new"].to_numpy() & ~catalog["_disc"].to_numpy() & (launch_day < 0))
    orders, items = [], []
    item_id = 1
    for n_order, (d, stream, cid) in enumerate(slots):
        order_id = 100001 + n_order
        date = days[d]
        online = stream == "Online"
        k = int(min(1 + rng.poisson(1.35 if online else 1.05), 8))
        chosen = rng.choice(len(sku), size=k, replace=False, p=month_p[(date.year, date.month)])
        cats = set()
        subtotal = 0
        is_bulk = n_order in bulk_at
        if is_bulk:
            chosen = np.concatenate([[rng.choice(bulk_pool)], chosen[chosen != chosen[0]]])[:k]
        for j, pi in enumerate(chosen):
            q = 1
            if maxq[pi] > 1 and rng.random() < 0.45:
                q = int(rng.integers(2, maxq[pi] + 1))
            if is_bulk and j == 0:
                q = int(rng.integers(20, 41))
            elif q > 1 and d >= price_cut and price_old[pi] != price_now[pi] and plant.random() < 0.6:
                q -= 1
            unit = int(price_old[pi] if d < price_cut else price_now[pi])
            subtotal += q * unit
            cats.add(cat_arr[pi])
            items.append((item_id, order_id, sku[pi], q, unit / 100))
            item_id += 1
        # promo codes
        code, disc = None, 0
        r = rng.random()
        if black_friday[d] and r < 0.60:
            code, disc = "SUMMIT20", 20
        elif promo[d] and r < 0.45:
            code, disc = "TRAIL15", 15
        elif cid not in first_seen and r < 0.30:
            code, disc = "WELCOME10", 10
        elif r < 0.03:
            code, disc = "GEARUP5", 5
        first_seen.add(cid)
        # the documented formula: round(items_subtotal * (1 - discount_pct / 100), 2) + shipping_cost
        after = int(round(float(np.round(subtotal / 100 * (1 - disc / 100), 2)) * 100))
        ship = 0
        if online:
            ship = 1495 if rng.random() < 0.07 else (0 if after >= 7500 else 695)
        # status
        p_ret = 0.035 + (0.10 if online and "Footwear" in cats else 0.0) + (0.05 if online and "Apparel" in cats else 0.0)
        p_can = 0.03 if online else 0.005
        r = rng.random()
        status = "returned" if r < p_ret else ("cancelled" if r < p_ret + p_can else "completed")
        if is_bulk:
            status = "completed"
        pay = (rng.choice(["card", "digital_wallet", "gift_card"], p=[0.62, 0.30, 0.08]) if online
               else rng.choice(["card", "cash", "digital_wallet", "gift_card"], p=[0.60, 0.22, 0.12, 0.06]))
        orders.append((order_id, cid, date, "online" if online else "store", stream, status, str(pay), code, disc,
                       ship / 100, (after + ship) / 100))
    orders = pd.DataFrame(orders, columns=["order_id", "customer_id", "order_date", "channel", "store", "status",
                                           "payment_method", "promo_code", "discount_pct", "shipping_cost", "order_total"])
    items = pd.DataFrame(items, columns=["order_item_id", "order_id", "sku", "quantity", "unit_price"])
    return orders, items


def finish_customers(cust: pd.DataFrame, orders: pd.DataFrame) -> pd.DataFrame:
    """Loyalty tier from completed spend, then drop the 'closed' accounts (orphan keys for merge lessons)."""
    rng = rng_for(7)
    spend = orders[orders["status"] == "completed"].groupby("customer_id")["order_total"].sum()
    s = cust["customer_id"].map(spend).fillna(0.0)
    cust = cust.assign(loyalty_tier=np.select([s >= 1500, s >= 600], ["Gold", "Silver"], "Basic"))
    buyers = cust.index[cust["customer_id"].isin(orders["customer_id"])].to_numpy()
    closed = rng.choice(buyers, N_CLOSED, replace=False)
    cust = cust.drop(index=closed).reset_index(drop=True)
    cols = ["customer_id", "first_name", "last_name", "email", "phone", "city", "state", "signup_date",
            "birth_year", "acquisition_channel", "loyalty_tier", "marketing_opt_in"]
    return cust[cols]


# --------------------------------------------------------------------------
# orders_messy.csv
# --------------------------------------------------------------------------
def make_messy(orders: pd.DataFrame):
    """The same orders, damaged on purpose. Returns (frame of strings, dict of exact flaw counts)."""
    rng = rng_for(8)
    n = len(orders)
    flaws: dict[str, int] = {}
    m = pd.DataFrame({
        "Order ID": orders["order_id"].astype(str),
        " Customer ID": orders["customer_id"].astype(str),
        "Order Date": orders["order_date"].dt.strftime("%Y-%m-%d"),
        "Channel": orders["channel"].astype(object),
        "Store": orders["store"].astype(object),
        "Status": orders["status"].astype(object),
        "Payment Method": orders["payment_method"].astype(object),
        "Promo Code": orders["promo_code"].astype(object).where(orders["promo_code"].notna(), ""),
        "Discount %": orders["discount_pct"].astype(str),
        "Shipping Cost": orders["shipping_cost"].map(lambda v: f"{v:.2f}"),
        "Order Total ": orders["order_total"].map(lambda v: f"{v:.2f}"),
    }).astype(object)

    def pick(k, pool=None):
        pool = np.arange(n) if pool is None else np.asarray(pool)
        return rng.choice(pool, k, replace=False)

    # 1 dates in three formats, some blank, some impossible
    us = pick(int(n * 0.09))
    m.loc[us, "Order Date"] = orders.loc[us, "order_date"].dt.strftime("%m/%d/%Y").to_numpy()
    rest = np.setdiff1d(np.arange(n), us)
    longf = pick(int(n * 0.04), rest)
    m.loc[longf, "Order Date"] = [f"{d:%B} {d.day}, {d.year}" for d in orders.loc[longf, "order_date"]]
    rest = np.setdiff1d(rest, longf)
    blank_d = pick(25, rest)
    m.loc[blank_d, "Order Date"] = ""
    rest = np.setdiff1d(rest, blank_d)
    bad_d = pick(6, rest)
    m.loc[bad_d, "Order Date"] = ["2024-02-30", "2025-02-29", "2024-13-01", "2025-04-31", "2024-06-31", "2025-00-15"]
    flaws.update(date_us_format=len(us), date_long_format=len(longf), date_blank=len(blank_d), date_impossible=len(bad_d))
    # 2 store names
    case = pick(int(n * 0.06))
    variants = [str.lower, str.upper, lambda s: " " + s, lambda s: s + " "]
    m.loc[case, "Store"] = [variants[rng.integers(4)](s) for s in m.loc[case, "Store"]]
    typo_map = {"Boise": "Bosie", "Flagstaff": "Flagstaf", "Asheville": "Ashville", "Bend": "Bend OR"}
    store_rows = np.setdiff1d(np.flatnonzero((orders["store"] != "Online").to_numpy()), case)
    typo = pick(160, store_rows)
    m.loc[typo, "Store"] = [typo_map[s] for s in m.loc[typo, "Store"]]
    online_rows = np.setdiff1d(np.flatnonzero((orders["store"] == "Online").to_numpy()), case)
    web = pick(int(len(online_rows) * 0.05), online_rows)
    m.loc[web, "Store"] = rng.choice(["Web", "website"], len(web))
    flaws.update(store_case_or_space=len(case), store_typo=len(typo), store_online_alias=len(web))
    # 3 status labels
    st = pick(int(n * 0.10))
    stat_map = {"completed": ["Completed", "COMPLETED", "complete", "completed "],
                "returned": ["Returned", "RETURNED", "returned ", "refunded"],
                "cancelled": ["Cancelled", "canceled", "CANCELLED", "cancelled "]}
    m.loc[st, "Status"] = [stat_map[s][rng.integers(4)] for s in m.loc[st, "Status"]]
    flaws["status_variant"] = len(st)
    # 4 customer id placeholders
    c1, c2, c3 = pick(60), None, None
    rest = np.setdiff1d(np.arange(n), c1)
    c2 = pick(40, rest)
    c3 = pick(25, np.setdiff1d(rest, c2))
    m.loc[c1, " Customer ID"] = ""
    m.loc[c2, " Customer ID"] = "N/A"
    m.loc[c3, " Customer ID"] = "unknown"
    flaws.update(customer_blank=60, customer_na_text=40, customer_unknown_text=25)
    # 5 order totals
    money = pick(int(n * 0.07))
    m.loc[money, "Order Total "] = [f"${v:,.2f}" for v in orders.loc[money, "order_total"]]
    rest = np.setdiff1d(np.arange(n), money)
    t_blank = pick(40, rest)
    m.loc[t_blank, "Order Total "] = ""
    rest = np.setdiff1d(rest, t_blank)
    returned = np.intersect1d(rest, np.flatnonzero((orders["status"] == "returned").to_numpy()))
    neg = pick(15, returned)
    m.loc[neg, "Order Total "] = [f"{-v:.2f}" for v in orders.loc[neg, "order_total"]]
    rest = np.setdiff1d(rest, neg)
    big = pick(8, rest)
    m.loc[big, "Order Total "] = [f"{v * 100:.2f}" for v in orders.loc[big, "order_total"]]
    flaws.update(total_currency_text=len(money), total_blank=40, total_negative=15, total_times_100=8)
    # 6 shipping
    zero_ship = np.flatnonzero((orders["shipping_cost"] == 0).to_numpy())
    free = pick(int(len(zero_ship) * 0.12), zero_ship)
    m.loc[free, "Shipping Cost"] = rng.choice(["free", "FREE"], len(free))
    s_blank = pick(30, np.setdiff1d(np.arange(n), free))
    m.loc[s_blank, "Shipping Cost"] = ""
    flaws.update(shipping_free_text=len(free), shipping_blank=30)
    # 7 payment method
    p_blank = pick(int(n * 0.02))
    m.loc[p_blank, "Payment Method"] = ""
    card_rows = np.setdiff1d(np.flatnonzero((orders["payment_method"] == "card").to_numpy()), p_blank)
    p_var = pick(int(len(card_rows) * 0.08), card_rows)
    m.loc[p_var, "Payment Method"] = rng.choice(["Card", "CARD", "credit card"], len(p_var))
    flaws.update(payment_blank=len(p_blank), payment_card_variant=len(p_var))
    # 8 promo codes
    has = np.flatnonzero(orders["promo_code"].notna().to_numpy())
    lower = pick(int(len(has) * 0.15), has)
    m.loc[lower, "Promo Code"] = [s.lower() for s in m.loc[lower, "Promo Code"]]
    none = np.flatnonzero(orders["promo_code"].isna().to_numpy())
    placeholder = pick(int(len(none) * 0.03), none)
    m.loc[placeholder, "Promo Code"] = rng.choice(["none", "-"], len(placeholder))
    flaws.update(promo_lowercase=len(lower), promo_placeholder=len(placeholder))
    # 9 exact duplicate rows, inserted right after a random later position
    dup = np.sort(pick(212))
    pieces = pd.concat([m, m.loc[dup]], ignore_index=True)
    order_key = np.concatenate([np.arange(n, dtype=float), dup + rng.integers(1, 40, len(dup)) + 0.5])
    m = pieces.iloc[np.argsort(order_key, kind="stable")].reset_index(drop=True)
    flaws["duplicate_rows"] = len(dup)
    # 10 a spreadsheet-style totals row at the bottom
    total_row = {c: "" for c in m.columns}
    total_row["Order ID"] = "TOTAL"
    total_row["Order Total "] = f"{orders['order_total'].sum():.2f}"
    m = pd.concat([m, pd.DataFrame([total_row])], ignore_index=True)
    flaws["total_row"] = 1
    flaws["rows_written"] = len(m)
    return m, flaws


# --------------------------------------------------------------------------
# Survey, A/B test, web sessions
# --------------------------------------------------------------------------
POS = ["Great selection and the staff really knew their gear.", "Fast shipping, boots fit perfectly.",
       "Love this store. Will be back before our next trip.", "Easy return, no questions asked.",
       "The staff helped me pick the right pack for a week-long trip.", "Good prices during the summer sale.",
       "Website was easy to use and my order came early.", "best gear shop in town!!",
       "Friendly people, clean store, found everything on my list."]
NEU = ["Fine overall. Nothing special.", "Prices are a little high but quality is good.",
       "Selection was ok, some sizes were out of stock.", "Shipping took about a week.",
       "Store was busy, had to wait for help.", "decent", "It was okay.   ",
       "Website search could be better but I found what I needed."]
NEG = ["Too expensive compared to other shops.", "Waited 20 minutes and nobody helped me.",
       "The website kept logging me out at checkout.", "Item arrived late and the box was damaged.",
       "OUT OF STOCK AGAIN. Very frustrating.", "Return took three weeks to refund.",
       "Could not find my size anywhere.", "checkout page is confusing on my phone",
       "Staff seemed rushed and not very helpful."]


def likert(rng, center, n):
    return np.clip(np.round(center + rng.normal(0, 0.62, n)), 1, 5)


def make_survey(cust: pd.DataFrame, orders: pd.DataFrame) -> pd.DataFrame:
    rng = rng_for(9)
    n = 1800
    o = orders[orders["customer_id"].isin(cust["customer_id"])]
    fav = o.groupby("customer_id")["store"].agg(lambda s: s.value_counts().index[0])
    pw = np.where(fav.to_numpy() == "Online", 1.0, 2.4)        # store shoppers were handed the survey more often
    ids = rng.choice(fav.index.to_numpy(), n, replace=False, p=pw / pw.sum())
    store = fav.loc[ids].to_numpy()
    by = cust.set_index("customer_id").loc[ids, "birth_year"].to_numpy()
    age = np.where(np.isnan(by), rng.integers(18, 78, n), 2025 - by)
    age_group = pd.cut(age, [0, 24, 34, 44, 54, 64, 200], labels=["18-24", "25-34", "35-44", "45-54", "55-64", "65+"]).astype(str)
    z = rng.normal(0, 1, n)
    online = store == "Online"
    staff_shift = np.where(store == "Asheville", -0.6, 0.0)
    sat_staff = likert(rng, 4.2 + 0.7 * z + staff_shift, n)
    sat_price = likert(rng, 3.4 + 0.6 * z + np.where(age_group == "18-24", -0.55, 0.0), n)
    sat_website = likert(rng, 3.7 + 0.6 * z + np.where(age_group == "65+", -0.65, 0.0), n)
    sat_selection = likert(rng, 4.1 + 0.65 * z, n)
    sat_overall = likert(rng, 4.0 + 0.8 * z + 0.4 * staff_shift, n)
    recommend = np.clip(np.round(7.6 + 1.9 * z + 0.9 * staff_shift + rng.normal(0, 1.1, n)), 0, 10)
    df = pd.DataFrame({
        "response_id": np.arange(1, n + 1), "customer_id": ids,
        "survey_date": pd.Timestamp("2025-10-01") + pd.to_timedelta(rng.integers(0, 46, n), unit="D"),
        "age_group": age_group, "store": store,
        "sat_overall": sat_overall, "sat_selection": sat_selection, "sat_price": sat_price,
        "sat_staff": sat_staff, "sat_website": sat_website, "recommend": recommend,
    })
    for col in ["sat_overall", "sat_selection", "sat_price", "sat_staff", "sat_website", "recommend"]:
        df.loc[rng.random(n) < 0.03, col] = np.nan                       # skipped at random
    df.loc[online & (rng.random(n) < 0.70), "sat_staff"] = np.nan        # online shoppers rarely met staff
    df.loc[~online & (rng.random(n) < 0.45), "sat_website"] = np.nan     # store shoppers often skip the website item
    comment = []
    for zi in z:
        if rng.random() < 0.55:
            comment.append(None)
        else:
            pool = POS if zi > 0.45 else (NEG if zi < -0.6 else NEU)
            comment.append(str(rng.choice(pool)))
    df["comment"] = comment
    return df.sort_values("survey_date", kind="stable").reset_index(drop=True).assign(response_id=lambda d: np.arange(1, n + 1))


def make_ab_test() -> pd.DataFrame:
    rng = rng_for(10)
    n = 28000
    variant = rng.choice(["control", "treatment"], n)
    device = rng.choice(["mobile", "desktop", "tablet"], n, p=[0.55, 0.38, 0.07])
    returning = rng.random(n) < 0.35
    base = pd.Series(device).map({"mobile": 0.060, "desktop": 0.105, "tablet": 0.080}).to_numpy()
    lift = pd.Series(device).map({"mobile": 1.28, "desktop": 1.01, "tablet": 1.10}).to_numpy()
    p = base * np.where(returning, 1.35, 1.0) * np.where(variant == "treatment", lift, 1.0)
    converted = rng.random(n) < p
    value = np.where(converted, np.round(rng.lognormal(4.35, 0.55, n), 2), np.nan)
    date = pd.Timestamp("2025-09-08") + pd.to_timedelta(np.sort(rng.integers(0, 28, n)), unit="D")
    return pd.DataFrame({"visitor_id": [f"V{i:06d}" for i in rng.permutation(n) + 1], "date": date,
                         "variant": variant, "device": device, "returning": returning,
                         "converted": converted.astype(int), "order_value": value})


def make_sessions(cust: pd.DataFrame, orders: pd.DataFrame, catalog: pd.DataFrame) -> dict:
    rng = rng_for(11)
    n = 2400
    nov = orders[(orders["channel"] == "online") & (orders["order_date"] >= "2025-11-01")
                 & (orders["order_date"] <= "2025-11-30")].reset_index(drop=True)
    buy_rows = rng.choice(len(nov), 230, replace=False)
    sellable = catalog[~catalog["_disc"] & ~catalog["_new"]]
    skus = sellable["sku"].to_numpy()
    price = dict(zip(sellable["sku"], sellable["price"]))
    cat_slug = {s: c.lower().replace(" ", "-") for s, c in zip(sellable["sku"], sellable["category"])}
    all_ids = cust["customer_id"].to_numpy()
    devices = [("mobile", "iOS", "Safari"), ("mobile", "Android", "Chrome"), ("desktop", "Windows", "Chrome"),
               ("desktop", "macOS", "Safari"), ("desktop", "Windows", "Edge"), ("desktop", "Linux", "Firefox"),
               ("tablet", "iPadOS", "Safari"), ("tablet", "Android", "Chrome")]
    dev_p = [0.30, 0.25, 0.18, 0.11, 0.06, 0.03, 0.05, 0.02]
    sources = [("search", "organic", None), ("search", "cpc", "fall_gear_search"), ("newsletter", "email", "november_trail_notes"),
               ("trailgram", "social", "winter_layers"), ("partner_blog", "referral", None)]
    sessions = []
    for i in range(n):
        buying = i < len(buy_rows)
        if buying:
            o = nov.iloc[buy_rows[i]]
            day, cid = o["order_date"], int(o["customer_id"])
        else:
            day = pd.Timestamp("2025-11-01") + pd.Timedelta(days=int(rng.integers(0, 30)))
            cid = int(rng.choice(all_ids)) if rng.random() < 0.28 else None
        start = day + pd.Timedelta(seconds=int(rng.integers(6 * 3600, 23 * 3600 + 1800)))
        dt, os_, br = devices[rng.choice(len(devices), p=dev_p)]
        k = int(min(1 + rng.poisson(4.2 if buying else 2.4), 14))
        pages, t = [], start
        viewed = []
        for j in range(k):
            r = rng.random()
            if j == 0 and r < 0.55:
                path = "/"
            elif r < 0.45:
                s = str(rng.choice(skus))
                viewed.append(s)
                path = f"/product/{s}"
            elif r < 0.75:
                path = f"/category/{cat_slug[str(rng.choice(skus))]}"
            elif r < 0.85:
                path = f"/search?q={str(rng.choice(['tent', 'rain jacket', 'boots', 'stove', 'headlamp', 'sleeping bag']))}".replace(" ", "+")
            else:
                path = str(rng.choice(["/sale", "/stores", "/help/returns", "/account"]))
            secs = int(rng.gamma(2.0, 22)) + 3
            pages.append({"path": path, "seconds": secs, "at": t.strftime("%Y-%m-%dT%H:%M:%S")})
            t += pd.Timedelta(seconds=secs)
        cart_items = []
        if buying or (viewed and rng.random() < 0.22):
            pool = viewed or [str(rng.choice(skus))]
            for s in dict.fromkeys(pool[: int(rng.integers(1, 4))]):
                cart_items.append({"sku": s, "qty": int(rng.integers(1, 3)), "unit_price": float(price[s])})
            for path in (["/cart", "/checkout"] if buying else ["/cart"]):
                secs = int(rng.gamma(2.0, 30)) + 5
                pages.append({"path": path, "seconds": secs, "at": t.strftime("%Y-%m-%dT%H:%M:%S")})
                t += pd.Timedelta(seconds=secs)
        sess = {"session_id": "", "customer_id": cid, "started_at": start.strftime("%Y-%m-%dT%H:%M:%S"),
                "device": {"type": dt, "os": os_, "browser": br}}
        if rng.random() < 0.62:                 # direct visits have no "utm" key at all
            src, med, camp = sources[rng.choice(len(sources), p=[0.40, 0.22, 0.16, 0.14, 0.08])]
            sess["utm"] = {"source": src, "medium": med}
            if camp:
                sess["utm"]["campaign"] = camp
        sess["pages"] = pages
        sess["cart"] = {"items": cart_items, "value": round(float(sum(c["qty"] * c["unit_price"] for c in cart_items)), 2)}
        sess["converted"] = bool(buying)
        sess["order_id"] = int(o["order_id"]) if buying else None
        sessions.append(sess)
    sessions.sort(key=lambda s: s["started_at"])
    for i, s in enumerate(sessions, 1):
        s["session_id"] = f"S{i:05d}"
    return {"site": "pinecrest-outfitters.example", "period": {"start": "2025-11-01", "end": "2025-11-30"},
            "note": "Fictional data written for the Python Data Science course.", "sessions": sessions}


# --------------------------------------------------------------------------
# Employees, inventory, the SQLite database, the workbook, the zip
# --------------------------------------------------------------------------
def make_employees() -> pd.DataFrame:
    rng = rng_for(12)
    rows = []
    used = set()

    def person():
        while True:
            f, l = str(rng.choice(FIRST)), str(rng.choice(LAST))
            if (f, l) not in used:
                used.add((f, l))
                return f, l

    def add(title, dept, location, manager, lo, hi, hired_after=2016):
        f, l = person()
        eid = 1 + len(rows)
        hire = pd.Timestamp(f"{hired_after}-01-01") + pd.Timedelta(days=int(rng.integers(0, (2025 - hired_after) * 365 + 200)))
        rows.append({"employee_id": eid, "first_name": f, "last_name": l, "title": title, "department": dept,
                     "location": location, "manager_id": manager, "hire_date": hire.strftime("%Y-%m-%d"),
                     "salary": int(round(rng.uniform(lo, hi), -2)),
                     "email": f"{f}.{l}@pinecrest.example.com".lower()})
        return eid

    ceo = add("Chief Executive Officer", "Executive", "HQ", None, 185000, 185000, 2016)
    vp_retail = add("VP of Retail", "Retail", "HQ", ceo, 138000, 146000)
    vp_ecom = add("VP of E-commerce", "E-commerce", "HQ", ceo, 136000, 144000)
    vp_ops = add("VP of Operations", "Operations", "HQ", ceo, 132000, 140000)
    west = add("Regional Manager, West", "Retail", "HQ", vp_retail, 98000, 104000)
    east = add("Regional Manager, East", "Retail", "HQ", vp_retail, 96000, 102000)
    for store, n_assoc in zip(STORES, [7, 5, 4, 5]):
        boss = east if store == "Asheville" else west
        mgr = add("Store Manager", "Retail", store, boss, 68000, 78000)
        leads = [add("Assistant Manager", "Retail", store, mgr, 50000, 56000, 2018) for _ in range(2)]
        for k in range(n_assoc):
            add("Sales Associate", "Retail", store, leads[k % 2], 33000, 41000, 2019)
    web = add("Web Manager", "E-commerce", "HQ", vp_ecom, 92000, 99000)
    add("Web Developer", "E-commerce", "HQ", web, 78000, 90000, 2019)
    add("Web Developer", "E-commerce", "HQ", web, 78000, 90000, 2019)
    mkt = add("Marketing Manager", "E-commerce", "HQ", vp_ecom, 86000, 93000)
    add("Email Marketing Specialist", "E-commerce", "HQ", mkt, 54000, 62000, 2019)
    analyst_lead = add("Analytics Lead", "Operations", "HQ", vp_ops, 94000, 102000)
    add("Data Analyst", "Operations", "HQ", analyst_lead, 66000, 76000, 2020)
    add("Data Analyst", "Operations", "HQ", analyst_lead, 66000, 76000, 2020)
    buyer = add("Head Buyer", "Operations", "HQ", vp_ops, 84000, 92000)
    add("Buyer", "Operations", "HQ", buyer, 58000, 66000, 2019)
    wh = add("Warehouse Manager", "Operations", "Warehouse", vp_ops, 70000, 77000)
    for _ in range(4):
        add("Warehouse Associate", "Operations", "Warehouse", wh, 36000, 43000, 2019)
    add("HR Manager", "Executive", "HQ", ceo, 82000, 90000)
    rows[0]["hire_date"] = "2016-02-01"          # the founder was there first
    return pd.DataFrame(rows)


def make_inventory(catalog: pd.DataFrame) -> pd.DataFrame:
    rng = rng_for(13)
    live = catalog[~catalog["_disc"]]
    rows = []
    for sku, pop, is_new in zip(live["sku"], live["_pop"], live["_new"]):
        sold_out = rng.random() < 0.05                 # nothing anywhere
        thin = rng.random() < 0.12                     # nearly gone
        for loc in STORES + ["Warehouse"]:
            scale = 9.0 if loc == "Warehouse" else 1.6
            qty = int(rng.poisson(scale * (2.5 + 4 * min(pop, 4))))
            if rng.random() < (0.02 if loc == "Warehouse" else 0.14):
                qty = 0
            if thin:
                qty = int(rng.integers(0, 3)) if loc != "Warehouse" else int(rng.integers(0, 9))
            if sold_out or (is_new and loc != "Warehouse"):
                qty = 0
            reorder = int(max(2, round(scale * (1.5 + 1.5 * min(pop, 4)))))
            counted = pd.Timestamp("2025-12-01") + pd.Timedelta(days=int(rng.integers(0, 31)))
            rows.append((sku, loc, qty, reorder, counted.strftime("%Y-%m-%d")))
    return pd.DataFrame(rows, columns=["sku", "location", "quantity", "reorder_point", "last_counted"])


SCHEMA = """
CREATE TABLE products (
    sku TEXT PRIMARY KEY, name TEXT NOT NULL, category TEXT NOT NULL, product_type TEXT NOT NULL,
    brand TEXT NOT NULL, price REAL NOT NULL, cost REAL NOT NULL, weight_kg REAL, color TEXT,
    rating REAL, review_count INTEGER NOT NULL, stock_status TEXT NOT NULL, launch_date TEXT NOT NULL);
CREATE TABLE customers (
    customer_id INTEGER PRIMARY KEY, first_name TEXT NOT NULL, last_name TEXT NOT NULL, email TEXT NOT NULL,
    phone TEXT, city TEXT NOT NULL, state TEXT NOT NULL, signup_date TEXT NOT NULL, birth_year INTEGER,
    acquisition_channel TEXT NOT NULL, loyalty_tier TEXT NOT NULL, marketing_opt_in INTEGER NOT NULL);
CREATE TABLE orders (
    order_id INTEGER PRIMARY KEY, customer_id INTEGER NOT NULL, order_date TEXT NOT NULL, channel TEXT NOT NULL,
    store TEXT NOT NULL, status TEXT NOT NULL, payment_method TEXT NOT NULL, promo_code TEXT,
    discount_pct INTEGER NOT NULL, shipping_cost REAL NOT NULL, order_total REAL NOT NULL);
CREATE TABLE order_items (
    order_item_id INTEGER PRIMARY KEY, order_id INTEGER NOT NULL REFERENCES orders(order_id),
    sku TEXT NOT NULL, quantity INTEGER NOT NULL, unit_price REAL NOT NULL);
CREATE TABLE employees (
    employee_id INTEGER PRIMARY KEY, first_name TEXT NOT NULL, last_name TEXT NOT NULL, title TEXT NOT NULL,
    department TEXT NOT NULL, location TEXT NOT NULL, manager_id INTEGER REFERENCES employees(employee_id),
    hire_date TEXT NOT NULL, salary INTEGER NOT NULL, email TEXT NOT NULL);
CREATE TABLE inventory (
    sku TEXT NOT NULL REFERENCES products(sku), location TEXT NOT NULL, quantity INTEGER NOT NULL,
    reorder_point INTEGER NOT NULL, last_counted TEXT NOT NULL, PRIMARY KEY (sku, location));
"""


def _records(df: pd.DataFrame):
    """Rows as plain Python values (None for missing) for sqlite3."""
    obj = df.astype(object).where(df.notna(), None)
    return [tuple(v.item() if hasattr(v, "item") else v for v in row) for row in obj.itertuples(index=False, name=None)]


def write_db(path: Path, tables: dict[str, pd.DataFrame]) -> None:
    if path.exists():
        path.unlink()
    con = sqlite3.connect(path)
    con.executescript(SCHEMA)
    for name, df in tables.items():
        df = df.copy()
        for col in df.columns:
            if pd.api.types.is_datetime64_any_dtype(df[col]):
                df[col] = df[col].dt.strftime("%Y-%m-%d")
            elif df[col].dtype == bool:
                df[col] = df[col].astype(int)
        if name == "customers":
            df["birth_year"] = df["birth_year"].astype("Int64")
        marks = ", ".join("?" * len(df.columns))
        con.executemany(f"INSERT INTO {name} ({', '.join(df.columns)}) VALUES ({marks})", _records(df))
    con.commit()
    con.execute("VACUUM")
    con.close()


def write_xlsx(path: Path, daily: pd.DataFrame, orders: pd.DataFrame) -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("  (openpyxl is not installed - skipped sales_2025.xlsx)")
        return False
    import datetime as dt
    d = daily[daily["date"].dt.year == 2025]
    monthly = d.pivot_table(index=d["date"].dt.month, columns="store", values="revenue", aggfunc="sum")[STORES].round(0)
    o = orders[(orders["order_date"].dt.year == 2025) & (orders["channel"] == "online") & (orders["status"] == "completed")]
    monthly["Online"] = o.groupby(o["order_date"].dt.month)["order_total"].sum().round(0)
    monthly.insert(0, "Month", pd.to_datetime(monthly.index.map(lambda m: f"2025-{m:02d}-01")).strftime("%b"))
    monthly = monthly.reset_index(drop=True)
    q = d.pivot_table(index="store", columns=d["date"].dt.quarter, values="revenue", aggfunc="sum")
    growth = {"Boise": 1.03, "Bend": 1.02, "Flagstaff": 0.94, "Asheville": 1.10}     # Asheville's targets were optimistic
    targets = [{"store": s, "quarter": f"Q{k}", "target_revenue": float(round(q.loc[s, k] * growth[s], -3))}
               for s in STORES for k in range(1, 5)]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        monthly.to_excel(xl, sheet_name="Monthly Sales", index=False, startrow=3)
        ws = xl.sheets["Monthly Sales"]
        ws["A1"] = "Pinecrest Outfitters - 2025 revenue by month (USD)"
        ws["A2"] = "Stores: point-of-sale totals. Online: completed orders."
        pd.DataFrame(targets).to_excel(xl, sheet_name="Targets", index=False)
        fixed = dt.datetime(2026, 1, 1)
        xl.book.properties.created = fixed
        xl.book.properties.modified = fixed
    # re-pack with fixed timestamps so the file is byte-identical on every run
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as out:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":      # openpyxl stamps "modified" with the current time
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-01-01T00:00:00Z\g<2>", data)
            out.writestr(zi, data)
    return True


ZIP_MEMBERS = ["README.md", "products.csv", "customers.csv", "orders.csv", "order_items.csv", "orders_messy.csv",
               "daily_sales.csv", "store_weather.csv", "survey.csv", "ab_test.csv", "web_sessions.json",
               "pinecrest.db", "sales_2025.xlsx", "make_data.py", "mock_api.py"]


def write_zip(out: Path) -> list[str]:
    """data/pinecrest-data.zip unpacks to a data/ folder next to the learner's notebooks."""
    missing = []
    with zipfile.ZipFile(out / "pinecrest-data.zip", "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for name in ZIP_MEMBERS:
            src = out / name
            if not src.exists():
                src = Path(__file__).parent / name
            if not src.exists():
                missing.append(name)
                continue
            zi = zipfile.ZipInfo(f"data/{name}", FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, src.read_bytes())
    return missing


# --------------------------------------------------------------------------
def build(out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    days = pd.date_range(START, END, freq="D")
    catalog = make_products()
    weather = make_weather(days)
    daily_full = make_daily_sales(days, weather)
    daily = punch_gaps(daily_full)
    cust_all = make_customers()
    orders, items = make_orders(cust_all, catalog, daily_full, days)
    customers = finish_customers(cust_all, orders)
    messy, flaws = make_messy(orders)
    survey = make_survey(customers, orders)
    ab = make_ab_test()
    sessions = make_sessions(customers, orders, catalog)
    employees = make_employees()
    inventory = make_inventory(catalog)

    total_stock = inventory.groupby("sku")["quantity"].sum()
    stock = catalog["sku"].map(total_stock)
    catalog["stock_status"] = np.select([stock == 0, stock < 20], ["out_of_stock", "low_stock"], "in_stock")
    products = catalog.loc[~catalog["_disc"], ["sku", "name", "category", "product_type", "brand", "price", "cost",
                                               "weight_kg", "color", "rating", "review_count", "stock_status",
                                               "launch_date"]].reset_index(drop=True)
    customers_csv = customers.assign(birth_year=customers["birth_year"].astype("Int64"))

    to_csv(products, out / "products.csv")
    to_csv(customers_csv, out / "customers.csv")
    to_csv(orders, out / "orders.csv")
    to_csv(items, out / "order_items.csv")
    messy.to_csv(out / "orders_messy.csv", index=False, lineterminator="\n", encoding="utf-8")
    to_csv(daily, out / "daily_sales.csv")
    to_csv(weather, out / "store_weather.csv")
    to_csv(survey.astype({c: "Int64" for c in ["sat_overall", "sat_selection", "sat_price", "sat_staff",
                                               "sat_website", "recommend"]}), out / "survey.csv")
    to_csv(ab, out / "ab_test.csv")
    (out / "web_sessions.json").write_text(json.dumps(sessions, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    write_db(out / "pinecrest.db", {"products": products, "customers": customers, "orders": orders,
                                    "order_items": items, "employees": employees, "inventory": inventory})
    write_xlsx(out / "sales_2025.xlsx", daily, orders)
    missing = write_zip(out)

    counts = {"products": len(products), "customers": len(customers), "orders": len(orders), "order_items": len(items),
              "orders_messy": len(messy), "daily_sales": len(daily), "store_weather": len(weather), "survey": len(survey),
              "ab_test": len(ab), "web_sessions": len(sessions["sessions"]), "employees": len(employees),
              "inventory": len(inventory)}
    print("Rows written:")
    for k, v in counts.items():
        print(f"  {k:<14}{v:>8,}")
    print("Deliberate flaws in orders_messy.csv:")
    for k, v in flaws.items():
        print(f"  {k:<24}{v:>7,}")
    if missing:
        print("Not in the zip (file not found):", ", ".join(missing))
    size = (out / "pinecrest-data.zip").stat().st_size / 1e6
    print(f"pinecrest-data.zip: {size:.2f} MB")
    return {"counts": counts, "flaws": flaws}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Generate the Pinecrest Outfitters course data.")
    ap.add_argument("--out", type=Path, default=Path(__file__).resolve().parent, help="output folder (default: this folder)")
    args = ap.parse_args()
    build(args.out)
